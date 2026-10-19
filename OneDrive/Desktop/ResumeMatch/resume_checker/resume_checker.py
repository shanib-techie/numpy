"""Resume text extraction + skill matching."""
import io
import re
import difflib
import zipfile

from docx import Document
from PyPDF2 import PdfReader

from .skills_db import ALIASES, ALL_SKILLS

GROUP_OF = {s.lower(): g for s, g in ALL_SKILLS}
SUPPORTED = (".pdf", ".docx")


def extract_text(filename, data):
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)
    raise ValueError("Only PDF and DOCX files are supported")


def normalize(text):
    text = text.lower()
    text = re.sub(r"[^a-z0-9+#.\s]", " ", text)      # keep c++, c#, .net, node.js
    return re.sub(r"\s+", " ", text).strip()


def _variants(skill):
    s = normalize(skill)
    out = {s, s.replace(" ", ""), s.replace(".", "")}
    out.update(normalize(a) for a in ALIASES.get(s, []))
    for key, alts in ALIASES.items():                 # reverse: "powerbi" -> "power bi"
        if s in {normalize(a) for a in alts} or s == key.replace(" ", ""):
            out.add(normalize(key))
    return {v for v in out if v}


def _pattern(variant):
    return re.compile(rf"(?<![a-z0-9+#]){re.escape(variant)}s?(?![a-z0-9+#])")


def match_skill(skill, text, tokens):
    """Return 'exact', 'fuzzy' or None."""
    for v in _variants(skill):
        if _pattern(v).search(text):
            return "exact"
    s = normalize(skill)
    if len(s) >= 6:                                   # typo tolerance for custom skills
        n = len(s.split())
        for i in range(len(tokens) - n + 1):
            if difflib.SequenceMatcher(None, s, " ".join(tokens[i:i + n])).ratio() >= 0.82:
                return "fuzzy"
    return None


def analyze_text(text, skills):
    norm = normalize(text)
    tokens = norm.split()
    seen, matched, missing = set(), [], []
    for raw in skills:
        skill = raw.strip()
        if not skill or skill.lower() in seen:
            continue
        seen.add(skill.lower())
        hit = match_skill(skill, norm, tokens)
        item = {"skill": skill, "group": GROUP_OF.get(skill.lower(), "Custom")}
        if hit:
            item["match"] = hit
            matched.append(item)
        else:
            missing.append(item)
    total = len(matched) + len(missing)
    pct = round(len(matched) / total * 100, 1) if total else 0.0
    return {
        "total": total, "matched": matched, "missing": missing,
        "match_percent": pct, "ats_score": pct, "verdict": verdict(pct),
        "suggestions": suggestions(missing, pct), "word_count": len(tokens),
    }


def verdict(pct):
    if pct >= 80: return "Strong match"
    if pct >= 60: return "Good match"
    if pct >= 40: return "Partial match"
    return "Weak match"


def suggestions(missing, pct):
    tips, by_group = [], {}
    for m in missing:
        by_group.setdefault(m["group"], []).append(m["skill"])
    for g, items in by_group.items():
        tips.append(f"Add {', '.join(items[:6])} ({g}) to your skills or projects if you have real experience with them.")
    if pct < 60:
        tips.append("Use the exact wording from the job description so ATS filters can find it.")
    if not missing:
        tips.append("Every skill was found. Back them up with measurable results in your experience section.")
    return tips


def iter_resumes(filename, data):
    """Yield (filename, bytes). A .zip upload is expanded into its resumes."""
    if filename.lower().endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                base = info.filename.rsplit("/", 1)[-1]
                if not info.is_dir() and base.lower().endswith(SUPPORTED) and not base.startswith("."):
                    yield base, z.read(info)
    else:
        yield filename, data


def analyze_file(filename, data, skills):
    result = {"filename": filename, "file_type": filename.rsplit(".", 1)[-1].upper()}
    try:
        text = extract_text(filename, data)
        if not text.strip():
            raise ValueError("No readable text found (scanned PDF?)")
        result.update(analyze_text(text, skills))
    except Exception as e:                           # one bad file must not stop the batch
        result["error"] = str(e)
    return result
